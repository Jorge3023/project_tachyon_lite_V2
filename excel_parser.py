from __future__ import annotations

import io
import re
import datetime
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import load_workbook, Workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import PatternFill

# 
#  CONFIGURACION
# 

SHEET_GRS = "GR'S"
SHEET_RM = "GRS (RM)"
SHEET_SQ00 = "SQ00"
SHEET_OUT = "SilverV2"

HEADER_ROW = 2          # fila de encabezados (1-indexado, igual que en Excel/VBA)
DATA_START_ROW = 3      # primera fila de datos

# Hojas fuente -> Status que se le asigna a cada una en la columna calculada "Status"
SOURCES: List[Tuple[str, str]] = [
    (SHEET_GRS, "GR"),
    (SHEET_RM, "GR"),
    (SHEET_SQ00, "OPO"),
]

# Encabezados de salida  EXACTAMENTE en este orden (igual que TargetHeaders() en VBA)
TARGET_HEADERS: List[str] = [
    "Unloading Point", "Customer name", "UNIQUE#", "MONTH", "Qtr", "PO Number",
    "Warning", "SO PRICE", "Vendor#", "Vendor Name", "Material#", "Quantity",
    "Price", "PO PRICE", "Per", "Delivery date", "SAP MTLS", "SAP REV", "SAP RMS",
    "GR Document", "TYPE", "REGION", "SITE", "Program", "PriceC", "ComCode",
    "Commodity", "Segment", "CC", "Cust Refresh", "P/n to be acct", "Customer Bought",
    "Customer Sold", "CEL P/N Bought", "CEL P/N Sold", "TRADER", "ROB NEW DEAL",
    "SPLIT PERCENT", "% RES", "HB months", "REP. MTLS", "REP. RMS", "RESERVES",
    "Region Indefinite", "RELEASE IN/FROM", "comments", "From (+) To (-) Reserves",
    "Splits (Regional (Site)) From original transaction", "SO Currency", "PO Currency",
    "Comments", "SO-CPN", "CC_1", "Sales Order", "Remarks", "MPN", "MFG",
    "Segment_2", "TAG", "SO Price", "Delta", "Status", "GR Date", "Category",
    "Month.", "RMS", "Year", "revenue", "Year-Month",
]

# Alias: nombre destino -> nombre real de la columna en el archivo origen
ALIAS_MAP: Dict[str, str] = {
    "quantity": "GR Quantity",
}

# Columnas que NUNCA se leen del origen: se calculan fila por fila
COMPUTED_COLS = {
    "month.", "rms", "year", "revenue", "year-month", "status",
}

_MONTH_STR_RE = re.compile(r"^(\d{4})\s+(\d{1,2})")
_EXCEL_EPOCH = datetime.datetime(1899, 12, 30)


# 
#  HELPERS
# 

def _norm(s: Any) -> str:
    return str(s).strip() if s is not None else ""


def _safe_num(v: Any) -> float:
    """Igual que SafeNum en VBA: si no es numérico, regresa 0."""
    if v is None:
        return 0.0
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    try:
        return float(str(v).strip().replace(",", ""))
    except (ValueError, TypeError):
        return 0.0


def _parse_year_month(raw: Any) -> Tuple[int, int]:
    """Igual que ParseYearMonth en VBA: intenta extraer (año, mes) de la columna MONTH."""
    if raw is None:
        return 0, 0

    if isinstance(raw, datetime.datetime):
        return raw.year, raw.month
    if isinstance(raw, datetime.date):
        return raw.year, raw.month

    # Números "sueltos" (sin formato de fecha) — VBA los trata como fecha serial
    if isinstance(raw, (int, float)) and not isinstance(raw, bool):
        try:
            dt = _EXCEL_EPOCH + datetime.timedelta(days=float(raw))
            if 1900 <= dt.year <= 2100:
                return dt.year, dt.month
        except (OverflowError, ValueError):
            pass
        return 0, 0

    s = _norm(raw)
    if not s:
        return 0, 0

    m = _MONTH_STR_RE.match(s)
    if m:
        yr = int(m.group(1))
        mo = int(m.group(2))
        if 1 <= mo <= 12:
            return yr, mo

    return 0, 0


def _sheet_exists(wb, name: str) -> bool:
    return name in wb.sheetnames


def _build_header_map(ws, header_row: int) -> Dict[str, int]:
    """Mapa nombre de columna (case-insensitive) -> índice de columna (0-indexado para iter_rows)."""
    header_map: Dict[str, int] = {}
    # Leemos de forma eficiente solo la fila que contiene los encabezados
    for row in ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True):
        for idx, val in enumerate(row):
            name = _norm(val)
            if name:
                key = name.lower()
                if key not in header_map:
                    header_map[key] = idx
        break  # Solo evaluamos la primera fila retornada (que es header_row)
    return header_map


# 
#  PROCESO PRINCIPAL
# 

def procesar_archivo(contenido: bytes, nombre: str) -> Dict[str, Any]:
    """
    Recibe los bytes del archivo Excel subido y regresa:
        {
          "resumen": {...},       # métricas para mostrar en la UI
          "excel_out": bytes,     # archivo .xlsx final (hoja SilverV2)
        }
    """
    if not nombre.lower().endswith((".xlsx", ".xls", ".xlsm")):
        raise ValueError(
            "El proceso Silver requiere un archivo Excel (.xlsx/.xlsm) que contenga "
            "las hojas \"GR'S\", \"GRS (RM)\" y/o \"SQ00\"."
        )

    try:
        wb = load_workbook(io.BytesIO(contenido), data_only=True, read_only=True)
    except Exception as e:
        raise ValueError(f"No se pudo abrir el archivo Excel: {e}")

    hojas_encontradas = [s for s, _ in SOURCES if _sheet_exists(wb, s)]
    if not hojas_encontradas:
        raise ValueError(
            "No se encontró ninguna de las hojas requeridas: "
            "\"GR'S\", \"GRS (RM)\", \"SQ00\". Verifica que no hayan sido renombradas."
        )

    idx_month = TARGET_HEADERS.index("MONTH")
    idx_sap_rms = TARGET_HEADERS.index("SAP RMS")
    idx_sap_rev = TARGET_HEADERS.index("SAP REV")

    filas_salida: List[List[Any]] = []
    conteo_por_status: Dict[str, int] = {}
    filas_omitidas_sin_fecha = 0
    hojas_procesadas = []
    hojas_ignoradas = []

    for sheet_name, status_val in SOURCES:
        if not _sheet_exists(wb, sheet_name):
            hojas_ignoradas.append(sheet_name)
            continue

        ws = wb[sheet_name]
        header_map = _build_header_map(ws, HEADER_ROW)

        # Resuelve, para cada columna destino, el índice de la columna origen (o None si es calculada / no existe)
        src_col_for: List[Optional[int]] = []
        for target_name in TARGET_HEADERS:
            key = target_name.lower()
            if key in COMPUTED_COLS:
                src_col_for.append(None)
            else:
                lookup_name = ALIAS_MAP.get(key, target_name)
                src_col_for.append(header_map.get(lookup_name.lower()))

        filas_leidas = 0
        filas_validas = 0

        # Iteramos las filas de forma secuencial y ultra rápida (values_only=True)
        for row_values in ws.iter_rows(min_row=DATA_START_ROW, values_only=True):
            filas_leidas += 1

            # Validar si toda la fila extraída viene vacía (final del archivo o fila en blanco)
            if all(v in (None, "") for v in row_values):
                continue

            month_col_idx = src_col_for[idx_month]
            month_raw = row_values[month_col_idx] if month_col_idx is not None and month_col_idx < len(row_values) else None

            yr, mo = _parse_year_month(month_raw)
            if yr == 0:
                filas_omitidas_sin_fecha += 1
                continue

            # Obtenemos los valores de SAP para calcular los reversos
            idx_rms_src = src_col_for[idx_sap_rms]
            idx_rev_src = src_col_for[idx_sap_rev]

            sap_rms = _safe_num(row_values[idx_rms_src] if idx_rms_src is not None and idx_rms_src < len(row_values) else None)
            sap_rev = _safe_num(row_values[idx_rev_src] if idx_rev_src is not None and idx_rev_src < len(row_values) else None)
            
            rms_val = sap_rms * -1
            revenue_val = sap_rev * -1
            year_month_txt = f"{yr}-{mo:02d}"

            fila_out: List[Any] = []
            for i, target_name in enumerate(TARGET_HEADERS):
                key = target_name.lower()
                if key == "month.":
                    fila_out.append(mo)
                elif key == "year":
                    fila_out.append(yr)
                elif key == "year-month":
                    fila_out.append(year_month_txt)
                elif key == "rms":
                    fila_out.append(rms_val)
                elif key == "revenue":
                    fila_out.append(revenue_val)
                elif key == "status":
                    fila_out.append(status_val)
                else:
                    col_idx = src_col_for[i]
                    if col_idx is not None and col_idx < len(row_values):
                        fila_out.append(row_values[col_idx])
                    else:
                        fila_out.append(None)

            filas_salida.append(fila_out)
            filas_validas += 1
            conteo_por_status[status_val] = conteo_por_status.get(status_val, 0) + 1

        hojas_procesadas.append({
            "hoja": sheet_name,
            "filas_leidas": filas_leidas,
            "filas_validas": filas_validas,
        })

    wb.close()

    # Ordena por "Year-Month" ascendente
    idx_year_month = TARGET_HEADERS.index("Year-Month")
    filas_salida.sort(key=lambda row: row[idx_year_month] or "")

    # Genera el archivo de salida 
    wb_out = Workbook()
    ws_out = wb_out.active
    ws_out.title = SHEET_OUT

    ws_out.append(TARGET_HEADERS)
    for row in filas_salida:
        ws_out.append(row)

    # Encabezado en negrita + autofiltro + freeze panes
    for cell in ws_out[1]:
        cell.font = cell.font.copy(bold=True)
    last_col_letter = get_column_letter(len(TARGET_HEADERS))
    ws_out.auto_filter.ref = f"A1:{last_col_letter}{len(filas_salida) + 1}"
    ws_out.freeze_panes = "A2"

    # Ajuste rapido de ancho de columnas
    for i, header in enumerate(TARGET_HEADERS, start=1):
        ws_out.column_dimensions[get_column_letter(i)].width = min(max(len(header) + 2, 10), 32)

    last_col_letter = get_column_letter(len(TARGET_HEADERS))
    ws_out.auto_filter.ref = f"A1:{last_col_letter}{len(filas_salida) + 1}"
    ws_out.freeze_panes = "A2"

    # Ajuste rapido de ancho de columnas
    for i, header in enumerate(TARGET_HEADERS, start=1):
        ws_out.column_dimensions[get_column_letter(i)].width = min(max(len(header) + 2, 10), 32)

    # ---> INICIO DEL NUEVO CÓDIGO DE COLORES <---
    # 1. Definimos los colores (Puedes cambiar los códigos Hexadecimales si prefieres otros tonos)
    color_verde = PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid")    # Verde pastel para GR
    color_amarillo = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid") # Amarillo pastel para OPO

    # 2. Encontramos en qué número de columna está "Status" (+1 porque openpyxl cuenta desde 1)
    col_status_idx = TARGET_HEADERS.index("Status") + 1

    # 3. Recorremos todas las filas (desde la fila 2 para saltar el encabezado)
    for row_idx in range(2, ws_out.max_row + 1):
        celda_status = ws_out.cell(row=row_idx, column=col_status_idx)
        
        if celda_status.value == "GR":
            celda_status.fill = color_verde
        elif celda_status.value == "OPO":
            celda_status.fill = color_amarillo

    # ---> FIN DEL CÓDIGO PARA DAR COLORES A GR Y OPO <---

    # Guardado del archivo
    buffer = io.BytesIO()
    wb_out.save(buffer)
    excel_bytes = buffer.getvalue()

    buffer = io.BytesIO()
    wb_out.save(buffer)
    excel_bytes = buffer.getvalue()

    #  Resumen para la UI 
    year_months = sorted({row[idx_year_month] for row in filas_salida if row[idx_year_month]})
    resumen = {
        "filas_totales": len(filas_salida),
        "columnas": len(TARGET_HEADERS),
        "por_status": conteo_por_status,
        "hojas_encontradas": hojas_encontradas,
        "hojas_ignoradas": hojas_ignoradas,
        "hojas_detalle": hojas_procesadas,
        "filas_omitidas_sin_fecha": filas_omitidas_sin_fecha,
        "rango_year_month": {
            "desde": year_months[0] if year_months else None,
            "hasta": year_months[-1] if year_months else None,
        },
        "hoja_salida": SHEET_OUT,
    }

    return {"resumen": resumen, "excel_out": excel_bytes}